import io
import unittest
from datetime import date, datetime, timezone
from pathlib import Path
from unittest.mock import patch

import pymupdf as fitz
from docx import Document
from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from backend.app.db.base import Base
from backend.app.models.policy import PolicyChunk, PolicyDocument
from backend.app.schemas.policy import PolicySummary, SearchFilters, SearchRequest
from backend.app.services.document_processing import (
    SourceSegment,
    extract_docx,
    extract_pdf,
    section_aware_chunks,
)
from backend.app.services.policy_service import ingest_policy, list_policies, search_policies


def make_pdf(text: str) -> bytes:
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), text)
    data = document.tobytes()
    document.close()
    return data


def make_docx() -> bytes:
    document = Document()
    document.add_heading("Coverage", level=1)
    document.add_paragraph("The policy covers emergency medical care.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Deductible"
    table.cell(0, 1).text = "$250"
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


class PolicyPipelineTests(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine(
            "sqlite://",
            connect_args={"check_same_thread": False},
            poolclass=StaticPool,
            execution_options={"schema_translate_map": {"insurance": None}},
        )
        Base.metadata.create_all(self.engine)
        self.session = Session(self.engine, expire_on_commit=False)

    def tearDown(self):
        self.session.close()
        self.engine.dispose()

    def test_pdf_extraction_keeps_page_number(self):
        segments = extract_pdf(make_pdf("Coverage\nEmergency medical care is included."))
        self.assertEqual(len(segments), 1)
        self.assertEqual(segments[0].page_number, 1)
        self.assertIn("Emergency medical care", segments[0].text)

    def test_docx_extraction_keeps_paragraphs_and_table_rows(self):
        segments = extract_docx(make_docx())
        combined = "\n".join(segment.text for segment in segments)
        self.assertIn("Coverage", combined)
        self.assertIn("emergency medical care", combined.lower())
        self.assertIn("Deductible | $250", combined)

    def test_chunking_retains_section_headings(self):
        chunks = section_aware_chunks(
            [
                SourceSegment(
                    text="Coverage\n" + ("Emergency care is included. " * 80),
                    page_number=2,
                    order=0,
                )
            ],
            chunk_size=300,
            chunk_overlap=40,
        )
        self.assertGreater(len(chunks), 1)
        self.assertTrue(all(chunk.section == "Coverage" for chunk in chunks))
        self.assertTrue(all(chunk.page_number == 2 for chunk in chunks))
        self.assertLessEqual(max(len(chunk.content) for chunk in chunks), 300)

    def test_ingestion_persists_source_metadata_and_embedded_chunks(self):
        vector = [1.0] + [0.0] * 1535
        with patch(
            "backend.app.services.policy_service.embed_texts",
            return_value=[vector],
        ), patch(
            "backend.app.services.policy_service.get_settings"
        ) as get_settings:
            settings = get_settings.return_value
            settings.max_upload_bytes = 1024 * 1024
            settings.upload_directory = Path("data/test-uploads")
            settings.chunk_size = 1000
            settings.chunk_overlap = 100
            settings.openai_embedding_model = "test-embeddings"
            policy = ingest_policy(
                self.session,
                data=make_pdf("Coverage\nEmergency medical care is included."),
                filename="health-policy.pdf",
                category="Health",
                carrier="Example Mutual",
                effective_date=date(2026, 1, 1),
            )
        self.assertEqual(policy.ingestion_status, "ready")
        self.assertEqual(policy.chunk_count, 1)
        self.assertEqual(policy.carrier, "Example Mutual")
        self.assertEqual(policy.effective_date, date(2026, 1, 1))
        self.assertEqual(len(self.session.query(PolicyChunk).all()), 1)
        Path(policy.storage_path).unlink(missing_ok=True)

    def test_revised_policy_increments_version_in_the_same_policy_series(self):
        vector = [1.0] + [0.0] * 1535
        with patch(
            "backend.app.services.policy_service.embed_texts",
            side_effect=lambda texts: [vector for _ in texts],
        ), patch(
            "backend.app.services.policy_service.get_settings"
        ) as get_settings:
            settings = get_settings.return_value
            settings.max_upload_bytes = 1024 * 1024
            settings.upload_directory = Path("data/test-uploads")
            settings.chunk_size = 1000
            settings.chunk_overlap = 100
            settings.openai_embedding_model = "test-embeddings"
            first = ingest_policy(
                self.session,
                data=make_pdf("Coverage\nEmergency care is included."),
                filename="health-policy.pdf",
                name="Health Policy",
                category="Health",
                carrier="Example Mutual",
            )
            second = ingest_policy(
                self.session,
                data=make_pdf("Coverage\nEmergency and urgent care are included."),
                filename="health-policy-revised.pdf",
                name="Health Policy",
                category="Health",
                carrier="Example Mutual",
            )
        self.assertEqual(first.policy_series_id, second.policy_series_id)
        self.assertEqual(first.version, 1)
        self.assertEqual(second.version, 2)
        Path(first.storage_path).unlink(missing_ok=True)
        Path(second.storage_path).unlink(missing_ok=True)

    def test_policy_status_reflects_expiry_effective_dates_and_archive(self):
        today = date.today()
        common = {
            "id": "policy",
            "policy_series_id": "series",
            "version": 1,
            "name": "Test",
            "filename": "test.pdf",
            "file_type": "pdf",
            "category": "Health",
            "carrier": None,
            "effective_date": None,
            "expires_date": None,
            "ingestion_status": "ready",
            "chunk_count": 1,
            "created_at": datetime.now(timezone.utc),
        }
        expired = PolicySummary.model_validate(
            {**common, "status": "active", "expires_date": date(today.year - 1, 12, 31)}
        )
        upcoming = PolicySummary.model_validate(
            {**common, "status": "active", "effective_date": date(today.year + 1, 1, 1)}
        )
        archived = PolicySummary.model_validate({**common, "status": "archived"})
        self.assertEqual(expired.status, "expired")
        self.assertEqual(upcoming.status, "upcoming")
        self.assertEqual(archived.status, "archived")

    def test_policy_listing_filters_effective_statuses(self):
        today = date.today()
        for policy_id, effective_date, expires_date, status in (
            ("active", None, None, "active"),
            ("expired", None, date(today.year - 1, 1, 1), "active"),
            ("upcoming", date(today.year + 1, 1, 1), None, "active"),
            ("archived", None, None, "archived"),
        ):
            self.session.add(
                PolicyDocument(
                    id=policy_id,
                    policy_series_id=f"series-{policy_id}",
                    version=1,
                    name=f"{policy_id.title()} policy",
                    filename=f"{policy_id}.pdf",
                    file_type="pdf",
                    storage_path="unused",
                    sha256=policy_id.ljust(64, "0"),
                    category="Health",
                    carrier=None,
                    effective_date=effective_date,
                    expires_date=expires_date,
                    status=status,
                    ingestion_status="ready",
                    chunk_count=0,
                    created_at=datetime.now(timezone.utc),
                    updated_at=datetime.now(timezone.utc),
                )
            )
        self.session.commit()
        for status, expected_id in (
            ("active", "active"),
            ("expired", "expired"),
            ("upcoming", "upcoming"),
            ("archived", "archived"),
        ):
            result = list_policies(self.session, status=status)
            self.assertEqual([item.id for item in result.items], [expected_id])

    def test_search_ranks_similar_chunks_and_respects_policy_filter(self):
        policy = PolicyDocument(
            id="policy-1",
            policy_series_id="series-1",
            version=1,
            name="Health Policy",
            filename="health.pdf",
            file_type="pdf",
            storage_path="unused",
            sha256="a" * 64,
            category="Health",
            carrier="Example",
            status="active",
            ingestion_status="ready",
            chunk_count=1,
            created_at=datetime.now(timezone.utc),
            updated_at=datetime.now(timezone.utc),
        )
        other = PolicyDocument(
            id="policy-2",
            policy_series_id="series-2",
            version=1,
            name="Auto Policy",
            filename="auto.pdf",
            file_type="pdf",
            storage_path="unused",
            sha256="b" * 64,
            category="Auto",
            carrier="Example",
            status="active",
            ingestion_status="ready",
            chunk_count=1,
            created_at=datetime.now(timezone.utc),
            updated_at=datetime.now(timezone.utc),
        )
        vector = [1.0] + [0.0] * 1535
        self.session.add_all(
            [
                policy,
                other,
                PolicyChunk(
                    id="chunk-1",
                    policy_id="policy-1",
                    chunk_index=0,
                    section="Coverage",
                    page_number=1,
                    content="Emergency care coverage.",
                    embedding=vector,
                    embedding_model="text-embedding-3-small",
                ),
                PolicyChunk(
                    id="chunk-2",
                    policy_id="policy-2",
                    chunk_index=0,
                    section="Liability",
                    page_number=1,
                    content="Auto liability coverage.",
                    embedding=[0.0, 1.0] + [0.0] * 1534,
                    embedding_model="text-embedding-3-small",
                ),
            ]
        )
        self.session.commit()
        with patch(
            "backend.app.services.policy_service.embed_texts",
            return_value=[vector],
        ):
            response = search_policies(
                self.session,
                SearchRequest(query="emergency care", top_k=5),
                SearchFilters(category="Health"),
            )
        self.assertEqual(response.total, 1)
        self.assertEqual(response.results[0].document_id, "policy-1")
        self.assertAlmostEqual(response.results[0].similarity_score, 1.0)

    def test_search_skips_embedding_call_when_no_indexed_policies_match(self):
        request = SearchRequest(query="emergency care", top_k=5)
        with patch("backend.app.services.policy_service.embed_texts") as embed:
            response = search_policies(self.session, request)
        embed.assert_not_called()
        self.assertEqual(response.results, [])
        self.assertEqual(response.total, 0)


if __name__ == "__main__":
    unittest.main()
